"""Generate the downloadable files for the `notion-kit` reward.

The reward is "Template Notion & Resume Starter Kit". Its Notion page carries the
tracker itself; this script builds the files attached to that page:

  dist/CV-Template-Indonesia.docx   one-page ATS-safe resume, Bahasa Indonesia
  dist/CV-Template-English.docx     the same resume in English
  dist/Tracker-Lamaran-Kerja.xlsx   Excel copy of the Notion tracker

Run from the repository root:  python docs/rewards/notion-resume-kit/build_kit.py

Resume rules (they decide whether an ATS can read the file at all): A4, one
column, no tables, text boxes, images, or page headers; contact details in the
body; standard font; dates aligned with a right tab stop. Sekolah Karir branding
never goes inside a resume — the file is sent to other companies.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
DIST = HERE / "dist"

NAVY = "07152D"
BLUE = "246BFD"
BLUE_TINT = "EBF1FF"
INK = "1A2540"

# ---------------------------------------------------------------------------
# Resumes
# ---------------------------------------------------------------------------

# Item kinds: ("para", text) · ("role", title, organisation, dates) ·
# ("bullet", text) · ("line", text) · ("kv", label, text)
RESUMES = {
    "CV-Template-Indonesia.docx": {
        "title": "CV - Nama Lengkap",
        "name": "NAMA LENGKAP",
        "headline": "Data Analyst  |  Fresh Graduate S1 Statistika",
        "contact": "Jakarta Selatan  |  0812-3456-7890  |  namakamu@email.com  |  linkedin.com/in/namakamu  |  github.com/namakamu",
        "sections": [
            ("RINGKASAN PROFIL", [
                ("para", "Lulusan S1 Statistika yang terbiasa mengolah data dengan SQL, Python, dan Excel, lalu menyajikannya "
                         "menjadi dashboard dan rekomendasi yang dipakai tim bisnis. Berpengalaman magang 6 bulan sebagai Data "
                         "Analyst dan menyelesaikan proyek analisis nyata di Side Hustle Arena. Mencari posisi Data Analyst "
                         "entry level di perusahaan teknologi atau ritel."),
            ]),
            ("PENGALAMAN", [
                ("role", "Data Analyst Intern", "PT Contoh Retail Nusantara, Jakarta", "Feb 2026 – Jul 2026"),
                ("bullet", "Membangun dashboard penjualan mingguan di Looker Studio untuk 40 cabang, memangkas waktu pembuatan "
                           "laporan dari 6 jam menjadi 30 menit per minggu."),
                ("bullet", "Menganalisis 1,2 juta baris data transaksi dengan SQL dan menemukan 5 produk bermargin terendah; "
                           "rekomendasi penyesuaian harganya dipakai tim kategori."),
                ("bullet", "Membuat skrip pembersihan data otomatis dengan Python (pandas) yang menurunkan kesalahan input dari "
                           "8% menjadi di bawah 1%."),
                ("role", "Asisten Praktikum Statistika", "Universitas Contoh, Bandung", "Agu 2024 – Jan 2025"),
                ("bullet", "Membimbing 60 mahasiswa per semester dalam praktikum R dan SPSS; 92% lulus ujian praktikum."),
            ]),
            ("PROYEK & BUKTI SKILL", [
                ("role", "Analisis Churn Pelanggan Aplikasi Langganan", "Side Hustle Arena, Sekolah Karir", "Sep 2026"),
                ("bullet", "Mengolah data 20.000 pelanggan untuk menemukan 3 faktor utama penyebab churn, lalu menyusun "
                           "rekomendasi retensi dalam laporan 5 halaman."),
                ("bullet", "Mendapat skor review 88/100 dan peringkat 2 dari 45 peserta divisi Data. Bukti: portfoliokamu.com/churn"),
                ("role", "Dashboard Tren Harga Bahan Pokok", "Proyek Pribadi", "Mar 2025"),
                ("bullet", "Mengumpulkan data harga 10 komoditas dari sumber publik dan memvisualisasikan tren 12 bulan di "
                           "Tableau Public; dilihat lebih dari 1.500 kali."),
            ]),
            ("PENDIDIKAN", [
                ("role", "S1 Statistika", "Universitas Contoh, Bandung", "2021 – 2025"),
                ("line", "IPK 3,62/4,00  |  Skripsi: Prediksi Permintaan Produk Ritel dengan Model ARIMA"),
            ]),
            ("ORGANISASI", [
                ("role", "Kepala Divisi Riset", "Himpunan Mahasiswa Statistika", "2023 – 2024"),
                ("bullet", "Memimpin 8 anggota menjalankan 4 survei kampus dengan total 1.200 responden."),
            ]),
            ("SKILL & SERTIFIKASI", [
                ("kv", "Hard skill", "SQL, Python (pandas, matplotlib), Excel (Pivot Table, XLOOKUP), Looker Studio, Tableau"),
                ("kv", "Soft skill", "Komunikasi data, presentasi ke stakeholder, kerja tim lintas fungsi"),
                ("kv", "Sertifikasi", "Google Data Analytics Professional Certificate (2025)"),
                ("kv", "Bahasa", "Indonesia (penutur asli), Inggris (profesional, TOEFL ITP 550)"),
            ]),
        ],
    },
    "CV-Template-English.docx": {
        "title": "CV - Full Name",
        "name": "FULL NAME",
        "headline": "Data Analyst  |  Statistics Graduate",
        "contact": "South Jakarta  |  +62 812-3456-7890  |  yourname@email.com  |  linkedin.com/in/yourname  |  github.com/yourname",
        "sections": [
            ("SUMMARY", [
                ("para", "Statistics graduate who turns raw data into dashboards and recommendations that business teams act on, "
                         "using SQL, Python, and Excel. Completed a 6-month Data Analyst internship and real-world analysis "
                         "projects on Side Hustle Arena. Seeking an entry-level Data Analyst role in technology or retail."),
            ]),
            ("EXPERIENCE", [
                ("role", "Data Analyst Intern", "PT Contoh Retail Nusantara, Jakarta", "Feb 2026 – Jul 2026"),
                ("bullet", "Built a weekly sales dashboard in Looker Studio covering 40 branches, cutting reporting time from "
                           "6 hours to 30 minutes per week."),
                ("bullet", "Analyzed 1.2 million transaction rows in SQL to identify the 5 lowest-margin products; the pricing "
                           "recommendations were adopted by the category team."),
                ("bullet", "Wrote a Python (pandas) data-cleaning script that cut input errors from 8% to under 1%."),
                ("role", "Statistics Lab Assistant", "Universitas Contoh, Bandung", "Aug 2024 – Jan 2025"),
                ("bullet", "Coached 60 students per semester in R and SPSS lab sessions; 92% passed the practical exam."),
            ]),
            ("PROJECTS & SKILL EVIDENCE", [
                ("role", "Customer Churn Analysis for a Subscription App", "Side Hustle Arena, Sekolah Karir", "Sep 2026"),
                ("bullet", "Analyzed data on 20,000 customers to identify the 3 main drivers of churn and wrote a 5-page "
                           "retention recommendation report."),
                ("bullet", "Scored 88/100 in project review and ranked 2nd of 45 participants in the Data division. "
                           "Evidence: yourportfolio.com/churn"),
                ("role", "Staple Food Price Trend Dashboard", "Personal Project", "Mar 2025"),
                ("bullet", "Collected prices for 10 commodities from public sources and visualized 12-month trends in Tableau "
                           "Public; viewed more than 1,500 times."),
            ]),
            ("EDUCATION", [
                ("role", "Bachelor of Statistics", "Universitas Contoh, Bandung", "2021 – 2025"),
                ("line", "GPA 3.62/4.00  |  Thesis: Retail Demand Forecasting with ARIMA Models"),
            ]),
            ("LEADERSHIP & ACTIVITIES", [
                ("role", "Head of Research Division", "Statistics Student Association", "2023 – 2024"),
                ("bullet", "Led 8 members in running 4 campus surveys with 1,200 respondents in total."),
            ]),
            ("SKILLS & CERTIFICATIONS", [
                ("kv", "Technical", "SQL, Python (pandas, matplotlib), Excel (Pivot Tables, XLOOKUP), Looker Studio, Tableau"),
                ("kv", "Soft skills", "Data storytelling, stakeholder presentations, cross-functional teamwork"),
                ("kv", "Certifications", "Google Data Analytics Professional Certificate (2025)"),
                ("kv", "Languages", "Indonesian (native), English (professional working proficiency, TOEFL ITP 550)"),
            ]),
        ],
    },
}

FONT = "Arial"
TEXT_WIDTH_CM = 21.0 - 2 * 1.9


def _set_run_font(run, size: float, bold: bool = False, color: str = INK, italic: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    # Without eastAsia/cs the run falls back to the theme font in some viewers.
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)


def _paragraph(doc, space_before: float = 0, space_after: float = 0, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = 1.08
    return paragraph


def _rule_below(paragraph, color: str = "9AA6B8") -> None:
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(key), value)
    border.append(bottom)
    # Word rejects a pPr whose children are out of schema order.
    paragraph._p.get_or_add_pPr().insert_element_before(
        border, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
        "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
        "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
        "w:rPr", "w:sectPr", "w:pPrChange",
    )


def build_resume(filename: str, spec: dict) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.6)
    section.left_margin = section.right_margin = Cm(1.9)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)

    props = doc.core_properties
    props.title = spec["title"]
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.keywords = ""

    name = _paragraph(doc)
    _set_run_font(name.add_run(spec["name"]), 18, bold=True)
    headline = _paragraph(doc, space_before=1)
    _set_run_font(headline.add_run(spec["headline"]), 10.5, bold=True, color="3A4A68")
    contact = _paragraph(doc, space_before=2)
    _set_run_font(contact.add_run(spec["contact"]), 9, color="3A4A68")

    for heading, items in spec["sections"]:
        head = _paragraph(doc, space_before=9, space_after=3)
        _rule_below(head)
        _set_run_font(head.add_run(heading), 10.5, bold=True)
        for kind, *parts in items:
            if kind == "para":
                p = _paragraph(doc, space_after=1)
                _set_run_font(p.add_run(parts[0]), 10)
            elif kind == "role":
                title, organisation, dates = parts
                p = _paragraph(doc, space_before=4, space_after=1)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
                _set_run_font(p.add_run(title), 10, bold=True)
                _set_run_font(p.add_run(f" — {organisation}"), 10)
                _set_run_font(p.add_run(f"\t{dates}"), 9.5, color="3A4A68")
            elif kind == "bullet":
                p = _paragraph(doc, space_after=1, style="List Bullet")
                p.paragraph_format.left_indent = Cm(0.55)
                p.paragraph_format.first_line_indent = Cm(-0.35)
                _set_run_font(p.add_run(parts[0]), 10)
            elif kind == "line":
                p = _paragraph(doc, space_after=1)
                _set_run_font(p.add_run(parts[0]), 10)
            elif kind == "kv":
                label, text = parts
                p = _paragraph(doc, space_after=1)
                _set_run_font(p.add_run(f"{label}: "), 10, bold=True)
                _set_run_font(p.add_run(text), 10)
            else:
                raise ValueError(f"Unknown resume item kind: {kind}")

    path = DIST / filename
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# Excel tracker
# ---------------------------------------------------------------------------

STATUSES = [
    # (name, fill, font colour, meaning)
    ("Wishlist", "F1F5F9", "475569", "Lowongan menarik yang belum kamu lamar."),
    ("Sudah Apply", "EBF1FF", "1A56D6", "Lamaran sudah terkirim, belum ada balasan."),
    ("Screening HR", "F3E8FF", "7E22CE", "HR menghubungi atau mengajak screening call."),
    # Colours follow the Notion tracker's select options, so both copies read alike.
    ("Tes / Case Study", "FCE7F3", "BE185D", "Diminta mengerjakan tes, psikotes, atau studi kasus."),
    ("Interview User", "FFEDD5", "C2410C", "Interview dengan calon atasan atau tim."),
    ("Offer", "FEF9C3", "A16207", "Menerima tawaran kerja, sedang dipertimbangkan."),
    ("Diterima", "16A34A", "FFFFFF", "Tawaran diterima. Selamat!"),
    ("Ditolak", "FEE2E2", "B91C1C", "Perusahaan menolak. Catat pelajarannya."),
    ("Tidak Ada Kabar", "EDE7E3", "78716C", "Lebih dari 3 minggu tanpa balasan setelah follow-up."),
]
PRIORITIES = ["Tinggi", "Sedang", "Rendah"]
SOURCES = ["LinkedIn", "JobStreet", "Glints", "Kalibrr", "Dealls", "Website Perusahaan", "Referral", "Lainnya"]
JOB_TYPES = ["Full-time", "Magang", "Kontrak", "Part-time", "Freelance"]

COLUMNS = [
    # (header, width)
    ("Posisi", 30), ("Perusahaan", 30), ("Status", 18), ("Prioritas", 11), ("Sumber", 19), ("Tipe", 12),
    ("Tanggal Apply", 15), ("Follow-up", 15), ("Link Lowongan", 30), ("Versi CV", 24), ("Kontak", 28),
    ("Ekspektasi Gaji", 17), ("Next Step", 42), ("Hari Sejak Apply", 16), ("Perlu Follow-up?", 22),
]

EXAMPLES = [
    ["Data Analyst Intern (contoh)", "PT Contoh Digital Indonesia", "Interview User", "Tinggi", "LinkedIn", "Magang",
     date(2026, 9, 1), date(2026, 9, 15), "https://www.linkedin.com/jobs/", "CV-ID-DataAnalyst-v2",
     "Kak Sari, Talent Acquisition", 4_500_000, "Latihan studi kasus SQL sebelum interview user"],
    ["Junior UI/UX Designer (contoh)", "Startup Contoh Kreatif", "Sudah Apply", "Sedang", "Glints", "Full-time",
     date(2026, 9, 5), date(2026, 9, 12), "https://glints.com/id", "CV-EN-UIUX-v1",
     "", 6_000_000, "Kirim email follow-up ke recruiter"],
    ["Product Analyst (contoh)", "PT Contoh Teknologi Nusantara", "Wishlist", "Tinggi", "Website Perusahaan", "Full-time",
     None, None, "", "", "Alumni kampus di tim Product", 7_000_000, "Minta referral lewat LinkedIn minggu ini"],
    ["Business Intelligence Trainee (contoh)", "Bank Contoh Sejahtera", "Ditolak", "Rendah", "JobStreet", "Kontrak",
     date(2026, 8, 20), None, "https://id.jobstreet.com", "CV-ID-DataAnalyst-v1",
     "", 5_000_000, "Minta feedback, lalu lamar posisi serupa"],
]

FIRST_ROW, LAST_ROW = 5, 204


def _banner(ws, last_col: str, title: str, subtitle: str) -> None:
    ws.merge_cells(f"A1:{last_col}1")
    ws.merge_cells(f"A2:{last_col}2")
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", start_color=NAVY, end_color=NAVY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Arial", size=10, color="FFFFFF")
    ws["A2"].fill = PatternFill("solid", start_color=BLUE, end_color=BLUE)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 22


def build_tracker() -> Path:
    wb = Workbook()
    wb.properties.creator = "Sekolah Karir"
    wb.properties.title = "Tracker Lamaran Kerja"

    ws = wb.active
    ws.title = "Tracker"
    last_col = "O"
    _banner(ws, last_col, "Tracker Lamaran Kerja",
            "Template dari Sekolah Karir, hadiah Side Hustle Arena. Hapus baris contoh, lalu isi lamaranmu sendiri.")

    for index, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=4, column=index, value=header)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[4].height = 22

    for offset, example in enumerate(EXAMPLES):
        row = FIRST_ROW + offset
        for index, value in enumerate(example, start=1):
            ws.cell(row=row, column=index, value=value)

    body_font = Font(name="Arial", size=10, color=INK)
    for row in range(FIRST_ROW, LAST_ROW + 1):
        ws.cell(row=row, column=14, value=f'=IF(G{row}="","",TODAY()-G{row})')
        ws.cell(row=row, column=15,
                value=f'=IF(OR(H{row}="",C{row}="Diterima",C{row}="Ditolak"),"",IF(H{row}<=TODAY(),"Waktunya follow-up",""))')
        for column in range(1, 16):
            cell = ws.cell(row=row, column=column)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=column in (1, 2, 11, 13))
        ws.cell(row=row, column=7).number_format = "dd mmm yyyy"
        ws.cell(row=row, column=8).number_format = "dd mmm yyyy"
        ws.cell(row=row, column=12).number_format = '"Rp"#,##0'
        ws.cell(row=row, column=14).number_format = "0"

    table = Table(displayName="TrackerLamaran", ref=f"A4:{last_col}{LAST_ROW}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    for values, column in (([s[0] for s in STATUSES], "C"), (PRIORITIES, "D"), (SOURCES, "E"), (JOB_TYPES, "F")):
        validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        validation.error = "Pilih salah satu dari daftar."
        validation.errorTitle = "Pilihan tidak tersedia"
        ws.add_data_validation(validation)
        validation.add(f"{column}{FIRST_ROW}:{column}{LAST_ROW}")

    status_range = f"C{FIRST_ROW}:C{LAST_ROW}"
    for name, fill, font_color, _ in STATUSES:
        ws.conditional_formatting.add(status_range, CellIsRule(
            operator="equal", formula=[f'"{name}"'],
            fill=PatternFill("solid", start_color=fill, end_color=fill),
            font=Font(color=font_color, bold=True)))
    ws.conditional_formatting.add(f"O{FIRST_ROW}:O{LAST_ROW}", CellIsRule(
        operator="equal", formula=['"Waktunya follow-up"'],
        fill=PatternFill("solid", start_color="FEF3C7", end_color="FEF3C7"), font=Font(color="B45309", bold=True)))

    ws.freeze_panes = "B5"
    ws.sheet_view.zoomScale = 100

    # Summary ----------------------------------------------------------------
    summary = wb.create_sheet("Ringkasan")
    _banner(summary, "H", "Ringkasan Lamaran", "Semua angka dihitung otomatis dari sheet Tracker.")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 12
    tracker_status = f"Tracker!$C${FIRST_ROW}:$C${LAST_ROW}"

    summary["A4"] = "Total lamaran di tracker"
    summary["B4"] = f"=COUNTA(Tracker!$A${FIRST_ROW}:$A${LAST_ROW})"
    summary["A6"], summary["B6"] = "Status", "Jumlah"
    for cell in (summary["A6"], summary["B6"]):
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=NAVY, end_color=NAVY)
    for offset, (name, fill, font_color, _) in enumerate(STATUSES):
        row = 7 + offset
        summary.cell(row=row, column=1, value=name).font = Font(name="Arial", size=10, bold=True, color=font_color)
        summary.cell(row=row, column=1).fill = PatternFill("solid", start_color=fill, end_color=fill)
        summary.cell(row=row, column=2, value=f"=COUNTIF({tracker_status},A{row})")
    # Rows: 7 Wishlist, 8 Sudah Apply, 9 Screening HR, 10 Tes, 11 Interview User, 12 Offer, 13 Diterima, 14 Ditolak, 15 Tidak Ada Kabar
    summary["A17"] = "Lamaran terkirim (tanpa Wishlist)"
    summary["B17"] = "=B4-B7"
    summary["A18"] = "Tingkat respons perusahaan"
    summary["B18"] = "=IFERROR((B17-B8-B15)/B17,0)"
    summary["A19"] = "Tingkat lolos sampai interview"
    summary["B19"] = "=IFERROR((B11+B12+B13)/B17,0)"
    summary["B18"].number_format = summary["B19"].number_format = "0%"
    for row in (4, 17, 18, 19):
        summary[f"A{row}"].font = Font(name="Arial", size=10, bold=True, color=INK)
        summary[f"B{row}"].font = Font(name="Arial", size=11, bold=True, color=BLUE)

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Lamaran per status"
    chart.legend = None
    chart.height, chart.width = 9, 16
    chart.add_data(Reference(summary, min_col=2, min_row=6, max_row=15), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=7, max_row=15))
    chart.y_axis.majorGridlines = None
    chart.x_axis.scaling.orientation = "maxMin"
    summary.add_chart(chart, "D4")

    # How-to -------------------------------------------------------------------
    guide = wb.create_sheet("Cara Pakai")
    _banner(guide, "C", "Cara Pakai Tracker", "Versi Excel dari template Notion Sekolah Karir. Bisa juga dibuka di Google Sheets.")
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 70
    steps = [
        ("Langkah 1", "Hapus 4 baris contoh di sheet Tracker (yang judulnya berakhiran \"(contoh)\")."),
        ("Langkah 2", "Setiap menemukan lowongan menarik, tambah satu baris dengan status Wishlist."),
        ("Langkah 3", "Setelah melamar, ganti status ke Sudah Apply, isi Tanggal Apply, dan pasang Follow-up 7 hari setelahnya."),
        ("Langkah 4", "Setiap ada kabar, geser status dan tulis Next Step. Kolom Perlu Follow-up menyala saat tanggalnya tiba."),
        ("Langkah 5", "Seminggu sekali buka sheet Ringkasan: lihat tingkat respons dan perbaiki CV kalau angkanya rendah."),
    ]
    row = 4
    for label, text in steps:
        guide.cell(row=row, column=1, value=label).font = Font(name="Arial", size=10, bold=True, color=BLUE)
        guide.cell(row=row, column=2, value=text).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    row += 1
    guide.cell(row=row, column=1, value="Arti setiap status").font = Font(name="Arial", size=11, bold=True, color=INK)
    row += 1
    for name, fill, font_color, meaning in STATUSES:
        label = guide.cell(row=row, column=1, value=name)
        label.font = Font(name="Arial", size=10, bold=True, color=font_color)
        label.fill = PatternFill("solid", start_color=fill, end_color=fill)
        guide.cell(row=row, column=2, value=meaning)
        row += 1
    row += 1
    guide.cell(row=row, column=1, value="Tips").font = Font(name="Arial", size=11, bold=True, color=INK)
    guide.cell(row=row + 1, column=2,
               value="Beri nama file CV sesuai versinya (misal CV-ID-DataAnalyst-v2) dan tulis nama itu di kolom Versi CV. "
                     "Saat dipanggil interview, kamu tahu persis CV mana yang dibaca perusahaan.").alignment = Alignment(wrap_text=True)

    wb.calculation.fullCalcOnLoad = True
    path = DIST / "Tracker-Lamaran-Kerja.xlsx"
    wb.save(path)
    return path


def main() -> None:
    DIST.mkdir(parents=True, exist_ok=True)
    for filename, spec in RESUMES.items():
        print("wrote", build_resume(filename, spec).relative_to(HERE))
    print("wrote", build_tracker().relative_to(HERE))


if __name__ == "__main__":
    main()
